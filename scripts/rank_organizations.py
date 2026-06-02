from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKBOOK_PATH = Path("data/Jody_Organizations_COMPLETE.xlsx")
SHEET_NAME = "Jody Organizations"
RANKED_OUTPUT_PATH = Path("data/Jody_Organizations_ranked.csv")
RESEARCH_OUTPUT_PATH = Path("data/Jody_Organizations_research_queue.csv")
DASHBOARD_OUTPUT_PATH = Path("data/Jody_Donation_Dashboard.csv")
RESEARCH_TOP20_OUTPUT_PATH = Path("data/Jody_Research_Top20.csv")


def normalize_header(value: str) -> str:
    normalized = value.lower().replace("©", "c")
    for token in (" ", "?", "%", ".", "(", ")", "/", "-", ",", "'"):
        normalized = normalized.replace(token, "_")
    while "__" in normalized:
        normalized = normalized.replace("__", "_")
    return normalized.strip("_")


def parse_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("%", "").replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_yes_no(value: Any) -> bool | None:
    if value in (None, ""):
        return None
    text = str(value).strip().lower()
    if text in ("y", "yes", "true", "1"):
        return True
    if text in ("n", "no", "false", "0"):
        return False
    return None


def clamp_score(score: float) -> float:
    return max(0.0, min(score, 100.0))


def calculate_mission_score(category: str, subcategory: str) -> float:
    category_text = (category or "").strip().lower()
    subcategory_text = (subcategory or "").strip().lower()

    base_map = {
        "animal": 85.0,
        "nature": 65.0,
        "veterans": 60.0,
        "alzheimers": 55.0,
        "science": 50.0,
        "political": 30.0,
    }
    score = base_map.get(category_text, 45.0)

    if "farm" in subcategory_text:
        score += 8.0
    if "wildlife" in subcategory_text:
        score += 6.0
    if "cats" in subcategory_text or "dogs" in subcategory_text:
        score += 3.0

    return clamp_score(score)


def calculate_efficiency_score(program_pct: Any, fundraising_pct: Any, admin_pct: Any) -> tuple[float, int]:
    score = 0.0
    present = 0

    program_value = parse_float(program_pct)
    if program_value is not None:
        present += 1
        if program_value >= 85:
            score += 45
        elif program_value >= 75:
            score += 35
        elif program_value >= 65:
            score += 25
        else:
            score += 10

    fundraising_value = parse_float(fundraising_pct)
    if fundraising_value is not None:
        present += 1
        if fundraising_value <= 12:
            score += 25
        elif fundraising_value <= 20:
            score += 20
        elif fundraising_value <= 30:
            score += 10
        else:
            score += 0

    admin_value = parse_float(admin_pct)
    if admin_value is not None:
        present += 1
        if admin_value <= 10:
            score += 20
        elif admin_value <= 15:
            score += 15
        elif admin_value <= 20:
            score += 8
        else:
            score += 0

    if present == 0:
        return 0.0, 0

    return clamp_score(score), present


def calculate_transparency_score(website: Any, ein: Any, is_501c3: Any, cn_rating: Any, notes: Any) -> tuple[float, int]:
    score = 0.0
    present = 0

    if website not in (None, ""):
        score += 20
        present += 1
    if ein not in (None, ""):
        score += 25
        present += 1

    is_verified = parse_yes_no(is_501c3)
    if is_verified is not None:
        present += 1
        if is_verified:
            score += 30

    rating_value = parse_float(cn_rating)
    if rating_value is not None:
        present += 1
        if rating_value >= 90:
            score += 25
        elif rating_value >= 80:
            score += 20
        elif rating_value >= 70:
            score += 14
        elif rating_value >= 60:
            score += 8
        else:
            score += 3

    if notes not in (None, ""):
        score += 5
        present += 1

    if present == 0:
        return 0.0, 0

    return clamp_score(score), present


def calculate_history_score(amount: Any) -> float:
    donation_amount = parse_float(amount) or 0.0
    if donation_amount >= 1000:
        return 100.0
    if donation_amount >= 500:
        return 80.0
    if donation_amount >= 250:
        return 60.0
    if donation_amount >= 100:
        return 40.0
    if donation_amount > 0:
        return 20.0
    return 0.0


def calculate_confidence_score(efficiency_fields_present: int, transparency_fields_present: int, duplicate_present: bool) -> float:
    present_count = efficiency_fields_present + transparency_fields_present + (1 if duplicate_present else 0)
    total_count = 3 + 5 + 1
    return round((present_count / total_count) * 100, 1)


def get_tier(final_score: float, confidence: float) -> str:
    if confidence < 45:
        return "Needs Research"
    if final_score >= 80:
        return "Tier 1"
    if final_score >= 65:
        return "Tier 2"
    return "Tier 3"


def get_next_step(missing_fields: list[str]) -> str:
    if not missing_fields:
        return "No action needed"

    next_field = missing_fields[0]
    return f"Fill {next_field}"


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    sheet = workbook[SHEET_NAME]
    rows = list(sheet.iter_rows(min_row=1, values_only=True))
    headers = [normalize_header(str(item)) for item in rows[0]]
    records: list[dict[str, Any]] = [dict(zip(headers, row)) for row in rows[1:]]

    ranked_rows: list[dict[str, Any]] = []
    research_rows: list[dict[str, Any]] = []
    dashboard_rows: list[dict[str, Any]] = []

    for record in records:
        mission_score = calculate_mission_score(
            str(record.get("category") or ""),
            str(record.get("subcategory") or ""),
        )
        efficiency_score, efficiency_fields_present = calculate_efficiency_score(
            record.get("program"),
            record.get("fundraising"),
            record.get("admin"),
        )
        transparency_score, transparency_fields_present = calculate_transparency_score(
            record.get("website"),
            record.get("ein_if_listed"),
            record.get("501cc3_verified_y_n"),
            record.get("charity_navigator_rating"),
            record.get("notes"),
        )
        history_score = calculate_history_score(record.get("approx_annual_donation_amount_if_known"))

        has_duplicate_value = parse_yes_no(record.get("duplicate_mission_y_n"))
        duplicate_penalty = 12.0 if has_duplicate_value is True else 0.0

        final_score = clamp_score(
            mission_score * 0.4
            + efficiency_score * 0.3
            + transparency_score * 0.25
            + history_score * 0.05
            - duplicate_penalty
        )

        confidence_score = calculate_confidence_score(
            efficiency_fields_present,
            transparency_fields_present,
            has_duplicate_value is not None,
        )

        critical_fields = {
            "Website": record.get("website"),
            "EIN": record.get("ein_if_listed"),
            "501c3 Verified": record.get("501cc3_verified_y_n"),
            "Charity Navigator Rating": record.get("charity_navigator_rating"),
            "Program %": record.get("program"),
            "Fundraising %": record.get("fundraising"),
            "Admin %": record.get("admin"),
        }
        missing_fields = [label for label, value in critical_fields.items() if value in (None, "")]
        donation_amount = parse_float(record.get("approx_annual_donation_amount_if_known")) or 0.0
        animal_multiplier = 1.2 if str(record.get("category") or "").lower() == "animal" else 1.0
        research_priority = round((len(missing_fields) + 1) * max(donation_amount, 1.0) * animal_multiplier, 2)

        ranked_row = {
            "Organization Name": record.get("organization_name"),
            "Category": record.get("category"),
            "Subcategory": record.get("subcategory"),
            "Approx Annual Donation": donation_amount,
            "Mission Score": round(mission_score, 1),
            "Efficiency Score": round(efficiency_score, 1),
            "Transparency Score": round(transparency_score, 1),
            "Legacy Fit Score": round(final_score, 1),
            "Confidence Score": confidence_score,
            "Tier": get_tier(final_score, confidence_score),
            "Missing Fields": ", ".join(missing_fields),
        }
        ranked_rows.append(ranked_row)

        research_row = {
            "Organization Name": record.get("organization_name"),
            "Category": record.get("category"),
            "Approx Annual Donation": donation_amount,
            "Research Priority": research_priority,
            "Missing Fields Count": len(missing_fields),
            "Missing Fields": ", ".join(missing_fields),
        }
        research_rows.append(research_row)

    ranked_rows.sort(
        key=lambda row: (
            row["Tier"] != "Tier 1",
            row["Tier"] != "Tier 2",
            -row["Legacy Fit Score"],
            row["Organization Name"] or "",
        )
    )
    research_rows.sort(key=lambda row: (-row["Research Priority"], row["Organization Name"] or ""))
    dashboard_source_rows = sorted(
        ranked_rows,
        key=lambda row: (-row["Legacy Fit Score"], -row["Approx Annual Donation"], row["Organization Name"] or ""),
    )

    for index, row in enumerate(dashboard_source_rows, start=1):
        dashboard_rows.append(
            {
                "Priority Rank": index,
                "Organization": row["Organization Name"],
                "Category": row["Category"],
                "Annual Donation": row["Approx Annual Donation"],
                "Legacy Score": row["Legacy Fit Score"],
                "Confidence": row["Confidence Score"],
                "Status": row["Tier"],
                "Next Step": get_next_step(
                    [field.strip() for field in str(row["Missing Fields"]).split(",") if field.strip()]
                ),
            }
        )

    with RANKED_OUTPUT_PATH.open("w", newline="", encoding="utf-8") as ranked_file:
        ranked_writer = csv.DictWriter(ranked_file, fieldnames=list(ranked_rows[0].keys()))
        ranked_writer.writeheader()
        ranked_writer.writerows(ranked_rows)

    with RESEARCH_OUTPUT_PATH.open("w", newline="", encoding="utf-8") as research_file:
        research_writer = csv.DictWriter(research_file, fieldnames=list(research_rows[0].keys()))
        research_writer.writeheader()
        research_writer.writerows(research_rows)

    with DASHBOARD_OUTPUT_PATH.open("w", newline="", encoding="utf-8") as dashboard_file:
        dashboard_writer = csv.DictWriter(dashboard_file, fieldnames=list(dashboard_rows[0].keys()))
        dashboard_writer.writeheader()
        dashboard_writer.writerows(dashboard_rows)

    top_20 = research_rows[:20]
    with RESEARCH_TOP20_OUTPUT_PATH.open("w", newline="", encoding="utf-8") as top_20_file:
        top_20_writer = csv.DictWriter(top_20_file, fieldnames=list(top_20[0].keys()))
        top_20_writer.writeheader()
        top_20_writer.writerows(top_20)

    print(f"Wrote {RANKED_OUTPUT_PATH}")
    print(f"Wrote {RESEARCH_OUTPUT_PATH}")
    print(f"Wrote {DASHBOARD_OUTPUT_PATH}")
    print(f"Wrote {RESEARCH_TOP20_OUTPUT_PATH}")


if __name__ == "__main__":
    main()
